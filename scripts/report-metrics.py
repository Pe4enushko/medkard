#!/usr/bin/env python3
"""Метрики прогона аудита по выгруженному отчёту и сравнение двух прогонов.

Зачем отдельно от e2e: у теста есть ожидаемый ответ, здесь его нет. Правки
промптов бинарным тестом не проверяются — «называй поле дословно» либо
сдвинуло долю замечаний с именем поля, либо нет, и увидеть это можно только
на прогоне. Вопрос тот же, что у scripts/deterministic-snapshot.py: не
«прошло или упало», а «что изменилось между двумя прогонами».

Чего скрипт не делает: не судит, стало лучше или хуже. Доли считаются грубо —
замечание считается назвавшим поле, если в его тексте встретилось имя поля
этой же карты. Формулировку «в поле, где указана температура» это не поймает,
и так и задумано: мерка должна быть тупой, иначе она измеряет саму себя.

Источник — выгруженный xlsx, а не БД, потому что отчёт есть всегда, а доступ к
`done_cards` бывает не у всех. Из отчёта карта восстанавливается не полностью:
хватает на тип визита и набор полей осмотра, не хватает на всё остальное.

Использование::

    python scripts/report-metrics.py snapshot report_Alenka_24-08-2026.xlsx -o after.json
    python scripts/report-metrics.py diff before.json after.json
"""
from __future__ import annotations

import argparse
import asyncio
import json
import re
import subprocess
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from audit.formal_structure.validator import FormalValidator  # noqa: E402

# Колонки ищутся по имени: состав отчёта менялся — «Источники КР» появилась
# между прогонами 23 и 24 августа и сдвинула проверку МКБ на позицию вправо.
_CARD = "Данные карты"
_INSPECTION = "Данные осмотра"
_SERVICES = "Услуги"
_DIAGNOSES = "Диагнозы"
_FORMAL = "Проверка по приказам МЗ"
_GUIDELINES = "Проверка по клин.рекоммендациям"
_ICD = "Проверка кодирования МКБ"
_SOURCES = "Источники КР"

_FLAG_RE = re.compile(r"\[([А-ЯЁ_0-9]{4,})\]")
_FORMAL_FINDING_RE = re.compile(
    r"\[[А-ЯЁ_0-9]{4,}\]\s*(.*?)(?:\n\s*\[Наблюдения\]:\s*(.*?))?(?=\n\s*\[Источник|\Z)", re.S
)
_GUIDELINE_FINDING_RE = re.compile(r"\[ЗАМЕЧАНИЕ\]\s*(.*?)(?=\n\s*\[ИСТОЧНИК|\n\s*─|\Z)", re.S)
_BLOCK_RE = re.compile(r"\n(?=\d+\.\n)")

# Значение, которое осталось от шаблона: голая единица измерения или прочерк.
_PLACEHOLDER_RE = re.compile(r"^(в мин\.?|мм рт\.? ст\.?|уд/мин|см|кг|-+|—|\.|n/a|нет данных)$", re.I)


def _read(path: str | Path) -> tuple[list[str], list[tuple]]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    if not rows:
        raise SystemExit(f"{path}: пустой лист")
    return [str(h or "") for h in rows[0]], rows[1:]


def _cell(headers: list[str], row: tuple, name: str) -> str:
    if name not in headers:
        return ""
    index = headers.index(name)
    return str(row[index] or "") if index < len(row) else ""


def _inspection_fields(cell: str) -> list[tuple[str, str]]:
    """Пары (Параметр, Значение) в том порядке, в каком они в отчёте."""
    out: list[tuple[str, str]] = []
    for block in _BLOCK_RE.split(cell):
        match = re.search(r"^\s*Значение:\s*(.*?)\n\s*Параметр:\s*(.+?)\s*$", block, re.S | re.M)
        if match:
            out.append((match.group(2).strip(), match.group(1).strip()))
    return out


def _visit(headers: list[str], row: tuple) -> dict[str, Any]:
    """Карта, восстановленная из отчёта настолько, чтобы её принял валидатор."""
    meta = _cell(headers, row, _CARD)

    def field(key: str) -> str | None:
        match = re.search(rf"^\s*{key}:\s*(.+?)\s*$", meta, re.M)
        return match.group(1) if match else None

    services = []
    for block in _BLOCK_RE.split(_cell(headers, row, _SERVICES)):
        code = re.search(r"^\s*КодЕГИСЗ:\s*(.+?)\s*$", block, re.M)
        name = re.search(r"^\s*Наименование:\s*(.+?)\s*$", block, re.M)
        if code or name:
            services.append({
                "КодЕГИСЗ": code.group(1) if code else "",
                "Наименование": name.group(1) if name else "",
            })

    diagnoses = [
        {"КодМКБ": match.group(1)}
        for block in _BLOCK_RE.split(_cell(headers, row, _DIAGNOSES))
        if (match := re.search(r"^\s*КодМКБ:\s*(.+?)\s*$", block, re.M))
    ]

    age = field("AGE")
    return {
        "Пациент": {"AGE": int(age) if age and age.isdigit() else age, "GENDER": field("GENDER")},
        "Врач": {"SPECIALIZATION": field("SPECIALIZATION")},
        "Прием": {"DATE": field("DATE"), "GUID": field("GUID")},
        "Услуги": services,
        "Диагнозы": diagnoses,
        "ДанныеОсмотра": [
            {"Параметр": param, "Значение": value}
            for param, value in _inspection_fields(_cell(headers, row, _INSPECTION))
        ],
    }


_WORD = r"[^\W_]"  # буква или цифра любого алфавита


def _names_a_field(text: str, labels: list[str]) -> bool:
    """Назвало ли замечание хоть одно поле этой карты дословно.

    Сверка по границам слова, а не по вхождению подстроки: имена полей бывают
    двух-трёхбуквенными («ЧСС», «ЧД», «Ф20»), и подстрока находила бы их внутри
    чужих слов. Отсечение по длине здесь не годится — оно выбрасывало бы ровно
    те аббревиатуры, ссылку на которые врач и ищет.

    Мерка намеренно грубая: «в поле, где указана температура» не засчитывается.
    """
    for label in labels:
        name = label.strip().strip(" .:;,-—")
        if len(name) < 2:
            continue
        if re.search(f"(?<!{_WORD}){re.escape(name)}(?!{_WORD})", text, re.I):
            return True
    return False


def _code_markers(headers: list[str], rows: list[tuple]) -> dict[str, Any]:
    """Признаки того, каким кодом снят отчёт.

    Сравнивать прогоны разных ревизий бессмысленно, а по имени файла это не
    видно: обе выгрузки называются одинаково и лежат рядом. Маркеры дешёвые и
    независимые — колонка источников КР появилась в одной ревизии, метка
    чекера МКБ сменилась в другой, а повтор одного флага в карте стал
    невозможен в третьей, когда правило начало давать один вердикт вместо
    массива замечаний.
    """
    icd = "".join(_cell(headers, row, _ICD) for row in rows)
    repeated = 0
    for row in rows:
        seen: dict[str, int] = {}
        for match in _FLAG_RE.finditer(_cell(headers, row, _FORMAL)):
            seen[match.group(1)] = seen.get(match.group(1), 0) + 1
        repeated += any(count > 1 for count in seen.values())
    return {
        "guideline_sources_column": _SOURCES in headers,
        "icd_label": (
            "recommendation" if "РЕКОМЕНДАЦИЯ ПО КОДУ МКБ" in icd
            else "error" if "ОШИБКА КОДИРОВАНИЯ МКБ" in icd
            else None
        ),
        "cards_with_a_repeated_flag": repeated,
    }


async def _metrics(headers: list[str], rows: list[tuple]) -> dict[str, Any]:
    validator = FormalValidator()
    flags: dict[str, int] = {}
    named = {"formal_issue": [0, 0], "formal_comment": [0, 0], "guideline_issue": [0, 0]}
    placeholder_cards = placeholder_flagged = 0
    incomplete = incomplete_mentioned = 0

    for row in rows:
        visit = _visit(headers, row)
        labels = [item["Параметр"] for item in visit["ДанныеОсмотра"]]
        formal = _cell(headers, row, _FORMAL)
        guidelines = _cell(headers, row, _GUIDELINES)

        for match in _FLAG_RE.finditer(formal):
            flags[match.group(1)] = flags.get(match.group(1), 0) + 1

        for issue, comment in _FORMAL_FINDING_RE.findall(formal):
            named["formal_issue"][1] += 1
            named["formal_issue"][0] += _names_a_field(issue, labels)
            if comment.strip():
                named["formal_comment"][1] += 1
                named["formal_comment"][0] += _names_a_field(comment, labels)
        for issue in _GUIDELINE_FINDING_RE.findall(guidelines):
            named["guideline_issue"][1] += 1
            named["guideline_issue"][0] += _names_a_field(issue, labels)

        if any(_PLACEHOLDER_RE.match(value.strip()) for _, value in
               _inspection_fields(_cell(headers, row, _INSPECTION))):
            placeholder_cards += 1
            placeholder_flagged += "ОБНАРУЖЕНЫ_ЗАГЛУШКИ" in formal

        unfilled = validator._check_missing_required_fields(  # noqa: SLF001
            visit, await validator.get_visit_types(visit)
        )
        if unfilled:
            incomplete += 1
            # сказала ли о том же сама модель — иначе неясно, что даёт код
            incomplete_mentioned += _names_a_field(formal + guidelines, [
                token.strip() for token in unfilled["issue"].split(":", 1)[1].split(",")
            ])

    return {
        "cards": len(rows),
        "flags": dict(sorted(flags.items(), key=lambda kv: (-kv[1], kv[0]))),
        "field_named": {key: {"named": value[0], "total": value[1]} for key, value in named.items()},
        "placeholders": {"cards": placeholder_cards, "flagged": placeholder_flagged},
        "missing_required": {"cards": incomplete, "also_named_by_model": incomplete_mentioned},
    }


def _revision() -> str:
    try:
        return subprocess.run(
            ["git", "-C", str(Path(__file__).resolve().parent.parent), "rev-parse", "--short", "HEAD"],
            capture_output=True, text=True, timeout=5,
        ).stdout.strip() or "unknown"
    except Exception:  # noqa: BLE001 — происхождение не обязано ронять снимок
        return "unknown"


async def _snapshot(path: str) -> dict[str, Any]:
    headers, rows = _read(path)
    out = await _metrics(headers, rows)
    out["_meta"] = {
        "report": Path(path).name,
        "snapshot_revision": _revision(),
        "columns": headers,
        "report_code_markers": _code_markers(headers, rows),
    }
    return out


def _share(part: dict[str, int], got_key: str, total_key: str) -> str:
    total, got = part.get(total_key, 0), part.get(got_key, 0)
    return f"{got}/{total}" + (f" ({100 * got // total}%)" if total else "")


def _diff(before: dict[str, Any], after: dict[str, Any]) -> int:
    meta_before, meta_after = before.pop("_meta", {}), after.pop("_meta", {})
    for label, meta in (("до   ", meta_before), ("после", meta_after)):
        markers = meta.get("report_code_markers", {})
        print(f"{label}: {meta.get('report', '?')}, маркеры кода: "
              f"источники КР={markers.get('guideline_sources_column')}, "
              f"метка МКБ={markers.get('icd_label')}, "
              f"карт с повтором флага={markers.get('cards_with_a_repeated_flag')}")
    if meta_before.get("report_code_markers") != meta_after.get("report_code_markers"):
        print("ВНИМАНИЕ: отчёты сняты разными ревизиями — расхождение метрик объясняется не только правкой\n")
    else:
        print()

    print(f"карт: {before.get('cards')} → {after.get('cards')}\n")

    print("замечаний, назвавших поле дословно:")
    for key in ("formal_issue", "formal_comment", "guideline_issue"):
        print(f"  {key:18} {_share(before['field_named'][key], 'named', 'total')}"
              f" → {_share(after['field_named'][key], 'named', 'total')}")

    print("\nзаглушки (карт с полем-заглушкой → из них с флагом):")
    print(f"  {_share(before['placeholders'], 'flagged', 'cards')}"
          f" → {_share(after['placeholders'], 'flagged', 'cards')}")

    print("\nнеполные записи (карт по коду → из них модель сказала о том же поле):")
    print(f"  {_share(before['missing_required'], 'also_named_by_model', 'cards')}"
          f" → {_share(after['missing_required'], 'also_named_by_model', 'cards')}")

    gained = {k: v for k, v in after["flags"].items() if k not in before["flags"]}
    lost = {k: v for k, v in before["flags"].items() if k not in after["flags"]}
    print("\nфлаги:")
    if gained:
        print(f"  появились: {gained}")
    if lost:
        print(f"  пропали:   {lost}")
    for flag in sorted(set(before["flags"]) & set(after["flags"])):
        if before["flags"][flag] != after["flags"][flag]:
            print(f"  {flag}: {before['flags'][flag]} → {after['flags'][flag]}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    commands = parser.add_subparsers(dest="command", required=True)

    take = commands.add_parser("snapshot", help="снять метрики с выгруженного отчёта")
    take.add_argument("report", help="xlsx отчёта аудита")
    take.add_argument("-o", "--out", help="файл снимка (по умолчанию — stdout)")

    compare = commands.add_parser("diff", help="сравнить два снимка")
    compare.add_argument("before")
    compare.add_argument("after")

    args = parser.parse_args()
    if args.command == "snapshot":
        snapshot = asyncio.run(_snapshot(args.report))
        text = json.dumps(snapshot, ensure_ascii=False, indent=2)
        if args.out:
            Path(args.out).write_text(text + "\n", encoding="utf-8")
            print(f"{snapshot['cards']} карт → {args.out}")
        else:
            print(text)
        return 0
    return _diff(
        json.loads(Path(args.before).read_text(encoding="utf-8")),
        json.loads(Path(args.after).read_text(encoding="utf-8")),
    )


if __name__ == "__main__":
    raise SystemExit(main())
